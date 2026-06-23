import csv
import io
import json
import re
import unicodedata
import zipfile
from urllib.parse import unquote
from xml.etree import ElementTree as ET

from database.dynamic_table_manager import delete_dynamic_row, list_dynamic_rows, upsert_dynamic_row
from database.expertise_manager import get_expertise, list_expertises, update_expertise
from services import cache_service
from services.runtime_context import clear_inventory_context, clear_skill_context


def _legacy_skill(e):
    try:
        persona = json.loads(e.get('persona_json') or '{}')
    except Exception:
        persona = {}
    return {
        'id': e['id'],
        'skill_id': str(persona.get('legacy_skill_id') or e['id']),
        'name': e.get('name') or '',
        'description': e.get('description') or e.get('job_title') or '',
        'data_table': e.get('data_table') or '',
        'data_fields_json': e.get('data_fields_json') or '[]',
    }


def _fields(e):
    try:
        arr = json.loads(e.get('data_fields_json') or '[]')
    except Exception:
        arr = []
    fields = []
    for idx, f in enumerate(arr if isinstance(arr, list) else []):
        if isinstance(f, str):
            key = f; label = f
        else:
            key = f.get('key') or f.get('field_key') or f.get('label') or f'field_{idx+1}'
            label = f.get('label') or f.get('field_label') or key
        fields.append({'field_key': key, 'field_label': label, 'field_type': 'text', 'required': bool(f.get('required')) if isinstance(f, dict) else False, 'display_order': idx})
    return fields


def _row_to_item(row, e):
    data = {'id': row.get('id'), 'content': row.get('content') or ''}
    fields = _fields(e) if e else []
    if fields:
        # The first/required field is the natural row id.
        id_key = _field_id_key(e) or fields[0].get('field_key')
        if id_key:
            data[id_key] = row.get('id')
        # Parse content lines of the form "Label: value" back into field values
        # so edit form/display can reuse existing data.
        label_to_key = {}
        for f in fields:
            label_to_key[str(f.get('field_label') or '').strip()] = f.get('field_key')
            label_to_key[str(f.get('field_key') or '').strip()] = f.get('field_key')
        for line in str(row.get('content') or '').splitlines():
            if ':' not in line:
                continue
            label, value = line.split(':', 1)
            key = label_to_key.get(label.strip())
            if key:
                data[key] = value.strip()
    return {'id': row.get('id'), 'data': data, 'content': row.get('content') or ''}


def _find_skill(selected_skill):
    selected_skill = unquote(selected_skill or '').strip()
    if selected_skill:
        e = get_expertise(selected_skill)
        if e:
            return e
    items = list_expertises()
    return items[0] if items else {}


def get_skill_data_page_context(selected_skill=None):
    skills = [_legacy_skill(e) for e in list_expertises()]
    e = _find_skill(selected_skill)
    rows = []
    if e and e.get('data_table'):
        rows = list_dynamic_rows(e['data_table'], limit=200)
    return {
        'skills': skills,
        'skill_name': str(e.get('id') or ''),
        'selected_skill': _legacy_skill(e) if e else {},
        'selected_skill_label': e.get('name') if e else '',
        'skill': _legacy_skill(e) if e else {},
        'fields': _fields(e) if e else [],
        'items': [_row_to_item(r, e) for r in rows],
        'items_total': len(rows),
    }


def api_get_fields(skill_name):
    e = get_expertise(unquote(skill_name))
    return {'success': True, 'fields': _fields(e) if e else []}, 200


def api_save_fields(skill_name, payload):
    e = get_expertise(unquote(skill_name))
    if not e:
        return {'success': False, 'errors': ['Không tìm thấy chuyên môn']}, 404
    payload = payload or {}
    raw_fields = payload.get('fields') or []
    if payload.get('quick_text') is not None:
        parts = [p.strip() for p in str(payload.get('quick_text') or '').replace(',', '\n').splitlines() if p.strip()]
        raw_fields = [{'key': p.lower().replace(' ', '_'), 'label': p, 'required': False} for p in parts]
    normalized = []
    for f in raw_fields:
        normalized.append({'key': f.get('key') or f.get('field_key') or f.get('label'), 'label': f.get('label') or f.get('field_label') or f.get('key'), 'required': bool(f.get('required'))})
    update_expertise(e['id'], {'data_fields_json': json.dumps(normalized, ensure_ascii=False)})
    clear_skill_context(str(e['id']))
    return {'success': True, 'fields': _fields(get_expertise(e['id']))}, 200


def api_get_items(skill_name, args):
    e = get_expertise(unquote(skill_name))
    if not e or not e.get('data_table'):
        return {'success': True, 'fields': [], 'items': [], 'total': 0, 'page': 1, 'page_size': 200}, 200
    rows = list_dynamic_rows(e['data_table'], query=(args or {}).get('query'), limit=int((args or {}).get('page_size', 200)))
    return {'success': True, 'fields': _fields(e), 'items': [_row_to_item(r, e) for r in rows], 'total': len(rows), 'page': 1, 'page_size': 200}, 200


def _data(payload):
    payload = payload or {}
    if isinstance(payload.get('data'), dict):
        data = dict(payload.get('data') or {})
        # UI sends the dynamic row id explicitly next to data. Keep it in the
        # data dict so legacy and new backends can resolve it consistently.
        if payload.get('id') not in (None, '') and data.get('id') in (None, ''):
            data['id'] = payload.get('id')
        return data
    return payload


def _field_id_key(e=None):
    fields = _fields(e) if e else []
    required = [f for f in fields if f.get('required') or f.get('is_required')]
    chosen = (required or fields or [{}])[0]
    return chosen.get('field_key') or chosen.get('key') or ''


def _row_id(d, e=None):
    """Return the row id for the dynamic data table.

    New data tables only have (id, content). The UI renders fields from
    expertises.data_fields_json; therefore the first declared field is the
    natural row id (for example: Mã biển, Mã sản phẩm, Mã dịch vụ).
    Do not require hard-coded legacy names such as ma_bien/ma_sp.
    """
    if not isinstance(d, dict):
        return ''

    # Explicit/common id keys first.
    for key in ('id', 'item_id', 'ma_bien', 'bien_so', 'ma_sp', 'product_code', 'sku', 'code', 'name'):
        value = d.get(key)
        if value not in (None, ''):
            return str(value).strip()

    # Then use the first field declared for this expertise. Prefer required.
    fields = _fields(e) if e else []
    ordered = [f for f in fields if f.get('required')] + [f for f in fields if not f.get('required')]
    for field in ordered:
        key = field.get('field_key')
        value = d.get(key)
        if value not in (None, ''):
            return str(value).strip()

    # Final fallback: first non-empty scalar in the payload.
    for value in d.values():
        if isinstance(value, (str, int, float)) and str(value).strip():
            return str(value).strip()
    return ''


def _content(d, e=None):
    # If the UI sends explicit raw content only, keep it.
    if 'content' in d and len([k for k in d.keys() if k not in {'id', 'item_id'}]) <= 1:
        return d.get('content') or ''

    # Build human-readable content using field labels from data_fields_json.
    labels = {f.get('field_key'): f.get('field_label') or f.get('field_key') for f in (_fields(e) if e else [])}
    lines = []
    for k, v in d.items():
        if k in {'item_id'}:
            continue
        # Keep the natural ID field in content too, because users expect
        # content to contain Mã biển/Mã sản phẩm as part of the record.
        if v is None or v == '':
            continue
        label = labels.get(k, k)
        lines.append(f'{label}: {v}')
    return '\n'.join(lines)


def _slug(value):
    """Normalize Vietnamese/English labels to comparable keys."""
    text = unicodedata.normalize('NFD', str(value or '').strip().lower())
    text = ''.join(ch for ch in text if unicodedata.category(ch) != 'Mn').replace('đ', 'd')
    text = re.sub(r'[^a-z0-9]+', '_', text).strip('_')
    return text


def _field_key_by_candidates(e, candidates):
    wanted = {_slug(c) for c in candidates}
    for field in _fields(e):
        key = field.get('field_key') or ''
        label = field.get('field_label') or ''
        if _slug(key) in wanted or _slug(label) in wanted:
            return key
    return ''


def _parse_money_value(value):
    if value is None or value == '':
        return 0
    if isinstance(value, (int, float)):
        return int(round(value))

    text = str(value).strip().lower()
    if not text:
        return 0

    text = text.replace('vnđ', '').replace('vnd', '').replace('đ', '')
    text = text.replace(' ', '').replace(',', '.')
    multiplier = 1
    if 'triệu' in text or 'trieu' in text or text.endswith('tr'):
        multiplier = 1_000_000
        text = text.replace('triệu', '').replace('trieu', '')
        text = re.sub(r'tr$', '', text)
    elif text.endswith('m'):
        multiplier = 1_000_000
        text = re.sub(r'm$', '', text)
    elif 'nghìn' in text or 'nghin' in text or text.endswith('k'):
        multiplier = 1_000
        text = text.replace('nghìn', '').replace('nghin', '')
        text = re.sub(r'k$', '', text)

    # 1.200.000 should be parsed as 1200000, while 1.2tr remains 1.2 * 1_000_000.
    if text.count('.') >= 2 or (multiplier == 1 and re.fullmatch(r'\d{1,3}(\.\d{3})+', text or '')):
        text = text.replace('.', '')
    try:
        return int(round(float(text) * multiplier))
    except Exception:
        digits = re.sub(r'\D', '', text)
        return int(digits) if digits else 0


def _format_cell_value(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_csv_rows(raw_bytes):
    text = raw_bytes.decode('utf-8-sig', errors='replace')
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
    except Exception:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [[_format_cell_value(cell) for cell in row] for row in reader]


def _read_xlsx_rows_stdlib(raw_bytes):
    """Small XLSX reader for simple sheets. Avoids hard dependency if openpyxl is absent."""
    ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
        shared_strings = []
        if 'xl/sharedStrings.xml' in zf.namelist():
            root = ET.fromstring(zf.read('xl/sharedStrings.xml'))
            for si in root.findall('x:si', ns):
                parts = [node.text or '' for node in si.findall('.//x:t', ns)]
                shared_strings.append(''.join(parts))

        workbook = ET.fromstring(zf.read('xl/workbook.xml'))
        first_sheet = workbook.find('x:sheets/x:sheet', ns)
        sheet_path = 'xl/worksheets/sheet1.xml'
        if first_sheet is not None:
            rel_id = first_sheet.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            if rel_id and 'xl/_rels/workbook.xml.rels' in zf.namelist():
                rels = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
                for rel in rels:
                    if rel.attrib.get('Id') == rel_id:
                        target = rel.attrib.get('Target') or 'worksheets/sheet1.xml'
                        sheet_path = 'xl/' + target.lstrip('/') if not target.startswith('xl/') else target
                        break

        root = ET.fromstring(zf.read(sheet_path))
        rows = []
        for row in root.findall('.//x:sheetData/x:row', ns):
            values = []
            last_col = 0
            for cell in row.findall('x:c', ns):
                ref = cell.attrib.get('r', '')
                letters = ''.join(ch for ch in ref if ch.isalpha())
                col_num = 0
                for ch in letters:
                    col_num = col_num * 26 + ord(ch.upper()) - 64
                while last_col + 1 < col_num:
                    values.append('')
                    last_col += 1

                cell_type = cell.attrib.get('t')
                value_node = cell.find('x:v', ns)
                inline_node = cell.find('x:is/x:t', ns)
                raw_value = value_node.text if value_node is not None else (inline_node.text if inline_node is not None else '')
                if cell_type == 's' and raw_value != '':
                    try:
                        raw_value = shared_strings[int(raw_value)]
                    except Exception:
                        pass
                values.append(_format_cell_value(raw_value))
                last_col = col_num or (last_col + 1)
            rows.append(values)
        return rows


def _read_uploaded_rows(uploaded_file):
    filename = (getattr(uploaded_file, 'filename', '') or '').lower()
    raw = uploaded_file.read()
    if not raw:
        return []
    if filename.endswith('.csv') or filename.endswith('.txt'):
        return _read_csv_rows(raw)

    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        return [[_format_cell_value(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    except ModuleNotFoundError:
        return _read_xlsx_rows_stdlib(raw)
    except Exception:
        # If openpyxl fails on a simple XLSX, try the dependency-free reader once.
        if filename.endswith(('.xlsx', '.xlsm')):
            return _read_xlsx_rows_stdlib(raw)
        raise


def _is_excel_header(row):
    slugs = {_slug(cell) for cell in row if str(cell or '').strip()}
    expected = {'bien_so', 'ma_bien', 'plate_number', 'gia', 'price', 'tinh', 'province'}
    return bool(slugs & expected)


def _value_from_row(row, header_map, index, candidates):
    if header_map:
        for candidate in candidates:
            idx = header_map.get(_slug(candidate))
            if idx is not None and idx < len(row):
                return row[idx]
    return row[index] if index < len(row) else ''


def api_import_excel_items(skill_name, uploaded_file):
    e = get_expertise(unquote(skill_name))
    if not e or not e.get('data_table'):
        return {'success': False, 'error': 'Chuyên môn chưa có bảng dữ liệu'}, 400
    if not uploaded_file:
        return {'success': False, 'error': 'Chưa chọn file Excel'}, 400

    filename = (uploaded_file.filename or '').lower()
    if not filename.endswith(('.xlsx', '.xlsm', '.csv', '.txt')):
        return {'success': False, 'error': 'Chỉ hỗ trợ file .xlsx, .xlsm hoặc .csv'}, 400

    try:
        rows = _read_uploaded_rows(uploaded_file)
    except Exception as exc:
        return {'success': False, 'error': f'Không đọc được file: {exc}'}, 400

    rows = [row for row in rows if any(str(cell or '').strip() for cell in row)]
    if not rows:
        return {'success': False, 'error': 'File không có dữ liệu'}, 400

    header_map = {}
    start_index = 0
    if _is_excel_header(rows[0]):
        header_map = {_slug(cell): idx for idx, cell in enumerate(rows[0]) if str(cell or '').strip()}
        start_index = 1

    plate_key = _field_key_by_candidates(e, ['plate_number', 'biển số', 'bien so', 'mã biển', 'ma bien', 'license plate']) or _field_id_key(e) or 'plate_number'
    price_key = _field_key_by_candidates(e, ['price', 'giá', 'gia', 'giá tiền', 'gia tien']) or 'price'
    province_key = _field_key_by_candidates(e, ['province', 'tỉnh', 'tinh', 'tỉnh thành', 'tinh thanh']) or 'province'
    vehicle_key = _field_key_by_candidates(e, ['vehicle_type', 'loại xe', 'loai xe']) or 'vehicle_type'
    status_key = _field_key_by_candidates(e, ['status', 'trạng thái', 'trang thai', 'tình trạng', 'tinh trang']) or 'status'

    imported = 0
    skipped = 0
    errors = []

    for excel_row_number, row in enumerate(rows[start_index:], start=start_index + 1):
        plate_number = _format_cell_value(_value_from_row(row, header_map, 0, ['biển số', 'bien so', 'mã biển', 'ma bien', 'plate_number', 'license plate']))
        price_value = _value_from_row(row, header_map, 1, ['giá', 'gia', 'price', 'giá tiền', 'gia tien'])
        province = _format_cell_value(_value_from_row(row, header_map, 2, ['tỉnh', 'tinh', 'province', 'tỉnh thành', 'tinh thanh']))

        if not plate_number:
            skipped += 1
            if len(errors) < 10:
                errors.append(f'Dòng {excel_row_number}: thiếu biển số')
            continue

        item_data = {
            plate_key: plate_number,
            price_key: _parse_money_value(price_value),
            province_key: province,
            vehicle_key: 'ô tô',
            status_key: 'còn',
        }
        try:
            upsert_dynamic_row(e['data_table'], plate_number, _content(item_data, e))
            imported += 1
        except Exception as exc:
            skipped += 1
            if len(errors) < 10:
                errors.append(f'Dòng {excel_row_number}: {exc}')

    clear_inventory_context(e['id'])
    return {
        'success': True,
        'imported': imported,
        'skipped': skipped,
        'errors': errors,
        'message': f'Đã import {imported} dòng. Bỏ qua {skipped} dòng.',
    }, 200


def api_create_item(skill_name, payload):
    e = get_expertise(unquote(skill_name))
    if not e or not e.get('data_table'):
        return {'success': False, 'errors': ['Chuyên môn chưa có bảng dữ liệu']}, 400
    d = _data(payload)
    rid = _row_id(d, e)
    if not rid:
        return {'success': False, 'errors': ['Thiếu ID dữ liệu']}, 400
    upsert_dynamic_row(e['data_table'], rid, _content(d, e))
    clear_inventory_context(e['id'])
    return {'success': True, 'item': {'id': rid}}, 200


def api_update_item(skill_name, item_id, payload):
    e = get_expertise(unquote(skill_name))
    if not e or not e.get('data_table'):
        return {'success': False, 'errors': ['Chuyên môn chưa có bảng dữ liệu']}, 400
    d = _data(payload)
    upsert_dynamic_row(e['data_table'], str(item_id), _content(d, e))
    clear_inventory_context(e['id'])
    return {'success': True}, 200


def api_delete_item(skill_name, item_id):
    e = get_expertise(unquote(skill_name))
    if not e or not e.get('data_table'):
        return {'success': False, 'error': 'Không tìm thấy dữ liệu'}, 404
    ok = delete_dynamic_row(e['data_table'], str(item_id))
    clear_inventory_context(e['id'])
    return {'success': ok}, 200 if ok else 404


def api_get_item(skill_name, item_id):
    e = get_expertise(unquote(skill_name))
    if not e or not e.get('data_table'):
        return {'success': False, 'error': 'Không tìm thấy dữ liệu'}, 404
    rows = [r for r in list_dynamic_rows(e['data_table'], limit=10000) if str(r['id']) == str(item_id)]
    if rows:
        r = rows[0]
        return {'success': True, 'item': _row_to_item(r, e)}, 200
    return {'success': False, 'error': 'Không tìm thấy dữ liệu'}, 404
